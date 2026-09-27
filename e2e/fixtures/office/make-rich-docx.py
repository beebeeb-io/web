#!/usr/bin/env python3
"""
make-rich-docx.py — task 1584: builds e2e/fixtures/office/rich-styled.docx,
a richly styled, multi-page Word document shaped like the file Guus opened on
his iPhone ("...-Resume-Styled.docx"): title + headings, bullet and numbered
lists, a table, an embedded PNG, and an EMBEDDED (obfuscated, ECMA-376
Part 1 §17.8.1) TrueType font that no device has installed.

The committed .docx is the fixture; this script is how it was made, so it can
be re-made. Requires python-docx + Pillow. Run from repos/web:

    python3 e2e/fixtures/office/make-rich-docx.py

The embedded font is Archivo Black (SIL Open Font License 1.1 — embedding is
permitted), read from --font (default: ~/Library/Fonts/ArchivoBlack-Regular.ttf).

The multi-chunk (> 4 MiB, the web client's CHUNK_SIZE) variant is NOT
committed — e2e/helpers/office-fixtures.ts derives it at test time by adding
incompressible padding to this file, so the repo does not carry a 5 MB blob.
"""
import argparse
import io
import os
import re
import sys
import uuid
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
EMBED_FONT_NAME = "Archivo Black"
# Fixed GUID so re-running the script produces the same obfuscation key.
FONT_KEY = "{5B1C2E3A-7D4F-4E61-9A2B-3C4D5E6F7A8B}"

# Text the e2e asserts on — keep in sync with e2e/1584-office-docx-rich.spec.ts.
TITLE = "Beebeeb Rich Fixture 1584"
HEADINGS = ["Experience", "Education", "Skills and Tools", "Projects", "References"]


def png_bytes() -> bytes:
    img = Image.new("RGB", (480, 240), (247, 243, 234))
    d = ImageDraw.Draw(img)
    for i in range(0, 480, 24):
        d.rectangle([i, 0, i + 12, 240], fill=(245, 184, 0) if (i // 24) % 2 else (36, 35, 32))
    d.ellipse([180, 60, 300, 180], fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def obfuscate(font: bytes, guid: str) -> bytes:
    hexs = re.sub(r"[^0-9A-Fa-f]", "", guid)
    # ECMA-376: key bytes are the GUID's hex digits read right-to-left in pairs.
    key = bytes(int(hexs[30 - 2 * i: 32 - 2 * i], 16) for i in range(16))
    out = bytearray(font)
    for i in range(32):
        out[i] ^= key[i % 16]
    return bytes(out)


def build_base(path: str) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)

    t = doc.add_paragraph()
    run = t.add_run(TITLE)
    run.font.name = EMBED_FONT_NAME
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x24, 0x23, 0x20)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    r = sub.add_run("Stored in Falkenstein. Encrypted before it left the device.")
    r.italic = True
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_picture(io.BytesIO(png_bytes()), width=Inches(4))

    for n, h in enumerate(HEADINGS):
        doc.add_heading(h, level=1)
        doc.add_heading(f"{h} detail", level=2)
        for p in range(6):
            para = doc.add_paragraph()
            para.add_run(f"{h} paragraph {p + 1}. ").bold = True
            para.add_run(
                "The quick brown fox jumps over the lazy dog while the encrypted "
                "vault keeps every byte sealed with a key only its owner holds. " * 3
            )
        for b in range(4):
            doc.add_paragraph(f"{h} bullet point {b + 1}", style="List Bullet")
        for b in range(3):
            doc.add_paragraph(f"{h} numbered step {b + 1}", style="List Number")
        if n == 1:
            table = doc.add_table(rows=4, cols=3)
            table.style = "Light Grid Accent 1"
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    cell.text = "Header" if ri == 0 else f"Cell {ri}.{ci}"
        doc.add_page_break()

    doc.save(path)


def embed_font(path: str, font_path: str) -> None:
    font = open(font_path, "rb").read()
    with zipfile.ZipFile(path) as zin:
        parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}

    ct = parts["[Content_Types].xml"].decode()
    if 'Extension="odttf"' not in ct:
        ct = ct.replace(
            "<Default ",
            '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ',
            1,
        )
    parts["[Content_Types].xml"] = ct.encode()

    ft = parts["word/fontTable.xml"].decode()
    if 'xmlns:r="' not in ft.split(">", 2)[1]:
        ft = ft.replace(
            "<w:fonts ",
            '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ',
            1,
        )
    font_el = (
        f'<w:font w:name="{EMBED_FONT_NAME}"><w:charset w:val="00"/><w:family w:val="auto"/>'
        f'<w:pitch w:val="variable"/><w:embedRegular r:id="rIdFont1" w:fontKey="{FONT_KEY}"/></w:font>'
    )
    ft = ft.replace("</w:fonts>", font_el + "</w:fonts>")
    parts["word/fontTable.xml"] = ft.encode()
    parts["word/_rels/fontTable.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdFont1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" '
        'Target="fonts/font1.odttf"/></Relationships>'
    ).encode()
    parts["word/fonts/font1.odttf"] = obfuscate(font, FONT_KEY)

    st = parts["word/settings.xml"].decode()
    if "embedTrueTypeFonts" not in st:
        st = re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/>", st, count=1)
    parts["word/settings.xml"] = st.encode()

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 27, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--font", default=os.path.expanduser("~/Library/Fonts/ArchivoBlack-Regular.ttf"))
    ap.add_argument("--out", default=os.path.join(HERE, "rich-styled.docx"))
    a = ap.parse_args()
    build_base(a.out)
    embed_font(a.out, a.font)
    print(f"wrote {a.out} ({os.path.getsize(a.out)} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
